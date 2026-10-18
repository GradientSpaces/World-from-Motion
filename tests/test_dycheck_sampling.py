# Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
from pathlib import Path
import sys

import numpy as np
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from wfm.dycheck_sampling import load_fixed_pose


def test_static_pose_preserves_world_to_camera_transform(tmp_path):
    pose = np.eye(4); pose[:3, 3] = [1., 2., 3.]
    path = tmp_path / "pose.npz"
    np.savez(path, poses=np.repeat(pose[None], 3, axis=0))
    np.testing.assert_array_equal(load_fixed_pose(path, 3), pose)


@pytest.mark.parametrize("failure", ["motion", "nan", "missing_frame", "homogeneous"])
def test_reject_incomparable_pose_archives(tmp_path, failure):
    poses = np.repeat(np.eye(4)[None], 3, axis=0)
    if failure == "motion":
        poses[1, 0, 3] = .1
    elif failure == "nan":
        poses[0, 0, 0] = np.nan
    elif failure == "missing_frame":
        poses = poses[:2]
    else:
        poses[:, 3, 3] = 0
    path = tmp_path / "poses.npz"; np.savez(path, poses=poses)
    with pytest.raises(ValueError):
        load_fixed_pose(path, 3)


def sample_fixture(tmp_path):
    import importlib.util
    import json
    from openpyxl import Workbook
    root = Path(__file__).resolve().parents[1]
    spec = importlib.util.spec_from_file_location("sampling_collector", root / "reproduce/sampling/collect.py")
    collector = importlib.util.module_from_spec(spec); spec.loader.exec_module(collector)
    config = json.loads((root / "config/dycheck_sampling.json").read_text())
    for method in ("ms", "wt"):
        run = tmp_path / f"{method}_wheel"; run.mkdir()
        source = tmp_path / f"source_{method}"; gt = source / "workspace/test_images"
        gt.mkdir(parents=True)
        names = [f"1_{i:05d}" for i in range(219)]
        for name in names:
            (gt / f"{name}.png").touch()
        for stage in ("align", "render", "generate", "score"):
            (run / f"{stage}-execution.json").write_text(json.dumps(dict(
                method=method, scene="wheel", stage=stage, status="complete", returncode=0,
                config=config, source_run=str(source), source_execution_sha256="a" * 64,
                elapsed_seconds=1.0)))
        (run / "source-unchanged.json").write_text('{"verified_initial_checkpoint_files": 3}')
        for folder in ("alignment", "scores"):
            out = run / folder; out.mkdir()
            book = Workbook(); sheet = book.active; sheet.append(["fn", *collector.table.STATIC])
            row = [20., .8, .2, 20., .8, .2, 20., .8, .2]
            for name in [*names, "AVE"]:
                sheet.append([name, *row])
            book.save(out / "tto_dycheck_metrics.xlsx"); book.close()
    return collector


def test_wheel_collection_does_not_claim_seven_scene_means(tmp_path):
    collector = sample_fixture(tmp_path)
    result = collector.collect(tmp_path, ["wheel"], tmp_path / "report.json")
    assert result["means"] == {} and len(result["scenes"]) == 2
    assert "partial" in result["scope"]


@pytest.mark.parametrize("failure", ["stage", "recipe", "frame", "initial_frame", "stage_identity", "source_identity"])
def test_collection_rejects_noncomparable_outputs(tmp_path, failure):
    import json
    from openpyxl import load_workbook
    collector = sample_fixture(tmp_path)
    run = tmp_path / "wt_wheel"
    if failure in ("stage", "recipe", "stage_identity", "source_identity"):
        p = run / "generate-execution.json"; data = json.loads(p.read_text())
        if failure == "stage":
            data["status"] = "failed"
        elif failure == "recipe":
            data["config"]["guidance_scale"] = 3.5
        elif failure == "stage_identity":
            data["stage"] = "align"
        else:
            data["source_execution_sha256"] = "b" * 64
        p.write_text(json.dumps(data))
    else:
        p = run / ("alignment" if failure == "initial_frame" else "scores") / "tto_dycheck_metrics.xlsx"
        book = load_workbook(p); book.active.cell(2, 1).value = "2_00000"; book.save(p); book.close()
    output = tmp_path / "report.json"
    with pytest.raises(ValueError):
        collector.collect(tmp_path, ["wheel"], output)
    assert not output.exists()
